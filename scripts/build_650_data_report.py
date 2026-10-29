# -*- coding: utf-8 -*-
"""build_650_data_report.py — data-only report for the 650 nm LIVE/DEAD dataset.

No interpretation, no summary, no recommendations: counting rules, excluded points, the
figures, and the numbers (per repeat, pooled, and the per-repeat Welch p values), plus
data-quality facts. All values are read from the CSVs written by plot_650_multirepeat.py
and from the raw ImageJ count tables.

Usage
  python build_650_data_report.py --figdir <figs> \
      --rep "repeat 1=<imagej_counts>" --rep "repeat 2=<imagej_counts_repeat2>" \
      --rep "repeat 3=<imagej_counts_repeat3>" \
      --exclusions <figs>/650nm_exclusions_20261029.csv \
      --out "F:/.../650nm_LIVE-DEAD_repeats1-3_data_report_20261029.pdf"
"""
from __future__ import annotations

import argparse
import os
import shutil

import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from scipy import stats

SCRATCH = 'C:/Users/PC/AppData/Local/hermes/cache/scratch'
TIMES = ['3h', '5h', '7h', '24h']
AREAS = ['c', '1', '2']
AREA_LABEL = {'c': 'Control', '1': 'LIPSS', '2': 'Nanopillar'}
SP_DIR = {'E. coli': 'E.coli', 'S. aureus': 'S.aureus'}


# ------------------------------------------------------------------ docx helpers
def add_bullets(doc, lines):
    for line in lines:
        doc.add_paragraph(line, style='List Bullet')


def add_note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(8.5)
    r.italic = True
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_figure(doc, path, width_cm, caption):
    if os.path.isfile(path):
        doc.add_picture(path, width=Cm(width_cm))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph('[missing figure: %s]' % os.path.basename(path))
    add_note(doc, caption)


def grid_table(doc, header, rows, caption=None, fontsize=7.5):
    if caption:
        doc.add_paragraph(caption, style='Heading 3')
    tb = doc.add_table(rows=1, cols=len(header))
    tb.style = 'Light Grid Accent 1'
    for i, h in enumerate(header):
        cell = tb.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(fontsize)
    for row in rows:
        cells = tb.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(fontsize)
    doc.add_paragraph()


def fmt_p(p):
    if p != p:
        return '-'
    if p < 0.001:
        return '<0.001'
    return '%.3f' % p


def sign(x):
    if not np.isfinite(x) or abs(x) <= 5:
        return '0'
    return '+' if x > 0 else '-'


def load_counts(dirpath, sp):
    p = os.path.join(dirpath, SP_DIR[sp], 'livedead_raw_counts.csv')
    if not os.path.isfile(p):
        return None
    df = pd.read_csv(p, encoding='utf-8-sig')
    df = df[df.area != '3'].copy()
    df['live'] = pd.to_numeric(df['live'], errors='coerce')
    df['dead'] = pd.to_numeric(df['dead'], errors='coerce')
    df['total'] = df['live'].fillna(0) + df['dead'].fillna(0)
    df['area'] = df['area'].astype(str)
    df['time'] = df['time'].astype(str)
    return df


def welch(raw, t, area, col='total'):
    if raw is None:
        return np.nan
    c = raw[(raw.time == t) & (raw.area == 'c')][col].astype(float)
    s = raw[(raw.time == t) & (raw.area == area)][col].astype(float)
    if len(c) < 2 or len(s) < 2:
        return np.nan
    return float(stats.ttest_ind(c, s, equal_var=False).pvalue)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--figdir', required=True)
    ap.add_argument('--rep', action='append', default=[])
    ap.add_argument('--exclusions', default='')
    ap.add_argument('--out', default='')
    ap.add_argument('--date', default='20261029')
    a = ap.parse_args()
    fig = a.figdir
    labels = [x.split('=', 1)[0].strip() for x in a.rep]
    dirs = {x.split('=', 1)[0].strip(): x.split('=', 1)[1].strip() for x in a.rep}

    R = {lb: pd.read_csv(os.path.join(fig, '650nm_livedead_%s_summary.csv'
                                      % lb.replace(' ', '')), encoding='utf-8-sig')
         for lb in labels}
    CB = pd.read_csv(os.path.join(fig, '650nm_livedead_combined_summary.csv'), encoding='utf-8-sig')
    AA = pd.read_csv(os.path.join(fig, '650nm_livedead_perrepeat_aa.csv'), encoding='utf-8-sig')
    AA['area'] = AA['area'].astype(str)
    for d in (CB,):
        d['area'] = d['area'].astype(str)
    for d in R.values():
        d['area'] = d['area'].astype(str)
    excl_path = a.exclusions or os.path.join(fig, '650nm_exclusions_%s.csv' % a.date)
    EX = (pd.read_csv(excl_path, encoding='utf-8-sig') if os.path.isfile(excl_path)
          else pd.DataFrame(columns=['level', 'repeat', 'species', 'time', 'area', 'issue',
                                     'detail', 'evidence']))
    raws = {lb: {sp: load_counts(dirs[lb], sp) for sp in ('E. coli', 'S. aureus')}
            for lb in labels}

    def rv(lb, sp, t, ar, col):
        df = R[lb]
        g = df[(df.species == sp) & (df.time == t) & (df.area == ar)]
        return float(g[col].iloc[0]) if len(g) else np.nan

    def aav(sp, t, ar, lb, live=False):
        col = '%s_AAlive%%' % lb if live else '%s_AA%%' % lb
        g = AA[(AA.species == sp) & (AA.time == t) & (AA.area == ar)]
        if not len(g) or col not in g.columns:
            return np.nan
        v = g[col].iloc[0]
        return float(v) if np.isfinite(v) else np.nan

    def pooled(sp, t, ar, live=False):
        g = AA[(AA.species == sp) & (AA.time == t) & (AA.area == ar)]
        if not len(g):
            return np.nan, np.nan, 0
        v = float(g['aa_live' if live else 'aa'].iloc[0])
        sd = float(g['aa_live_sd' if live else 'aa_sd'].iloc[0])
        n = int(g['n_aa'].iloc[0])
        return v, sd, n

    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)

    doc.add_heading('650 nm laser-textured stainless steel - bacterial LIVE/DEAD assay', level=0)
    doc.add_paragraph('Data report: %d batches (%s), generated %s. '
                      'Contains counting rules, the excluded data points, the figures and the '
                      'measured values; no interpretation is given.'
                      % (len(labels), ', '.join(labels), a.date))

    # ---------------------------------------------------------------- 1. reduction rules
    doc.add_heading('1. Data reduction rules (as applied)', level=1)
    add_bullets(doc, [
        'Coupon: one 650 nm textured stainless-steel coupon per condition per batch, three areas '
        'per coupon - c = untreated control, 1 = LIPSS, 2 = nanopillar.',
        'Four coupons rotate through the four incubation points, one place per batch: batch 1 = '
        '1,2,3,4; batch 2 = 2,3,4,1; batch 3 = 3,4,1,2 (coupon identity per incubation point).',
        'Organisms: E. coli, S. aureus. Incubation points: 3, 5, 7, 24 h. Five fields per area '
        'and incubation point, each imaged in a green (live) and a red (dead) channel. One batch '
        '= 8 folders x 30 images = 240 images.',
        'Counting: ImageJ 1.54 headless, adaptive threshold (histogram background mode + 40), '
        'Analyze Particles size 3-500 px, 40x field area 46 946.10 um2.',
        'Density = count / 46 946.10 um2 x 1e8 cells/cm2. Viability = pooled live / (live + dead) '
        'over the fields of a condition. Anti-adhesion AA = (N_control - N_surface) / N_control x '
        '100, computed for total cells (live + dead) and for live cells separately, against the '
        'control of the same batch and incubation point.',
        'Per-batch AA carries a delta-method SD propagated from the field means. Pooled AA = '
        'arithmetic mean of the per-batch AA values; its SD is the SD across batches (given only '
        'when at least two batches contribute). No value is borrowed or estimated across batches: '
        'an excluded point is not replaced, and every pooled number is a mean of measurements '
        'with their own controls.',
        'Welch t-test (control vs treated, n = 5 fields per batch), uncorrected; p values are '
        'reported per batch.',
    ])

    # ---------------------------------------------------------------- 2. exclusions
    doc.add_heading('2. Excluded data points', level=1)
    if len(EX):
        rows, reasons = [], []
        n_grp = n_aa = 0
        for _, r in EX.iterrows():
            area = '' if pd.isna(r['area']) else str(r['area'])
            try:
                area = str(int(float(area))) if area not in ('', 'c') else area
            except ValueError:
                pass
            n_grp += 3 if not area else 1
            n_aa += 2 if not area else 1
            rows.append([r['repeat'], r['time'], r['species'],
                         'whole incubation point' if not area else AREA_LABEL.get(area, area),
                         r['issue']])
            reasons.append('%s, %s, %s: %s (%s)'
                           % (r['repeat'], r['time'], r['species'], r['detail'], r['evidence']))
        grid_table(doc, ['Batch', 'Incubation point', 'Organism', 'Area', 'Issue'], rows,
                   caption='Table 1 - data points removed from the dataset.')
        add_bullets(doc, reasons)
        add_bullets(doc, [
            'Rule order: reported experimental problems remove a point outright; then an '
            'objective screen for repeat-level AA outliers (a batch is flagged only when the '
            'other two agree within 25 percentage points and the flagged one lies more than 100 '
            'percentage points from their mean; with two disagreeing batches nothing is removed); '
            'then a field-level screen (>3x the median of the other four fields), which flagged '
            'four fields - all retained, because their objects are single cells of normal size '
            '(no >500 px objects) and removing them would change the group mean.',
            '%d of the %d condition groups are removed, i.e. %d of the %d AA values are absent. '
            'Where the control itself is removed (batch 2, S. aureus 7 h and 24 h) no AA exists '
            'for that incubation point; the affected slots are drawn as excluded in the figures.'
            % (n_grp, 2 * 4 * 3 * len(labels), n_aa, 2 * 4 * 2 * len(labels)),
        ])
    else:
        doc.add_paragraph('No exclusions were applied.')

    # ---------------------------------------------------------------- 3. figures
    doc.add_heading('3. Figures', level=1)
    doc.add_paragraph('Each figure: left = live/dead density (cells/cm2) with pooled viability '
                      'over each bar; right = anti-adhesion (%) for total cells (solid) and live '
                      'cells (pale). Positive = fewer cells than the control of the same batch '
                      'and incubation point. Grey hatched bars tagged EXCLUDED, and grey x markers '
                      'in the combined figure, are the removed measurements: plotted with their '
                      'values, not used in any statistic.')
    figs = [(lb, '650nm_livedead_%s_AA.png' % lb.replace(' ', '')) for lb in labels]
    figs.append(('combined', '650nm_livedead_combined_%s_AA.png'
                 % ''.join(lb.replace('repeat', 'r').replace(' ', '') for lb in labels)))
    for k, (lb, fn) in enumerate(figs, start=1):
        is_comb = lb == 'combined'
        doc.add_heading('Figure %d - %s' % (k, 'all batches combined' if is_comb else lb),
                        level=2)
        if is_comb:
            cap = ('Figure %d. Combined over %s: bars = mean of the per-batch AA values, error '
                   'bars = SD across batches, empty circle / square / triangle = the per-batch '
                   'values, grey x = values removed from the dataset.'
                   % (k, ', '.join(labels)))
        else:
            cap = ('Figure %d. %s: left = live (green) + dead (red) density of the 5 fields, %% '
                   'over the bar = pooled viability; right = anti-adhesion of total (solid) and '
                   'live (pale) cells, error bars = delta-method SD of the field means.'
                   % (k, lb))
            mine = EX[EX['repeat'] == lb] if len(EX) else EX
            if len(mine):
                bits = []
                for _, r in mine.iterrows():
                    area = '' if pd.isna(r['area']) else str(r['area'])
                    try:
                        area = str(int(float(area))) if area not in ('', 'c') else area
                    except ValueError:
                        pass
                    bits.append('%s %s %s'
                                % (r['time'], r['species'],
                                   'whole incubation point' if not area
                                   else AREA_LABEL.get(area, area)))
                cap += ' Excluded in this batch: ' + '; '.join(bits) + '.'
        add_figure(doc, os.path.join(fig, fn), 16.4, cap)

    # ---------------------------------------------------------------- 4. tables
    doc.add_heading('4. Tables', level=1)

    for sp in ('E. coli', 'S. aureus'):
        hdr = ['Time']
        for lb in labels:
            hdr += ['%s c' % lb.replace('repeat ', 'R'), 'LIPSS', 'Nano']
        rows = []
        for t in TIMES:
            row = [t.replace('h', ' h')]
            for lb in labels:
                for ar in AREAS:
                    v = rv(lb, sp, t, ar, 'total_cm2') / 1e5
                    row.append('-' if not np.isfinite(v) else '%.1f' % v)
            rows.append(row)
        grid_table(doc, hdr, rows,
                   caption='Table 4.%d - %s: total cells (live + dead), x1e5 cells/cm2; "-" = '
                   'excluded.' % (1 if sp == 'E. coli' else 2, sp), fontsize=8)

        rows = []
        for t in TIMES:
            row = [t.replace('h', ' h')]
            for lb in labels:
                for ar in AREAS:
                    v = rv(lb, sp, t, ar, 'viability')
                    row.append('-' if not np.isfinite(v) else '%.0f%%' % v)
            rows.append(row)
        grid_table(doc, hdr, rows,
                   caption='Table 4.%d - %s: pooled viability, live / (live + dead); "-" = '
                   'excluded.' % (3 if sp == 'E. coli' else 4, sp), fontsize=8)

    for live, tag, num in ((False, 'total cells', 5), (True, 'live cells', 6)):
        rows = []
        for sp in ('E. coli', 'S. aureus'):
            for ar, nm in (('1', 'LIPSS'), ('2', 'Nanopillar')):
                for t in TIMES:
                    vals = [aav(sp, t, ar, lb, live=live) for lb in labels]
                    mv, sdv, n = pooled(sp, t, ar, live=live)
                    sg = {sign(v) for v in vals if np.isfinite(v)} - {'0'}
                    if n == 0:
                        desc = '-'
                    elif n < 2:
                        desc = 'n.a. (n = %d)' % n
                    else:
                        desc = 'same sign' if len(sg) <= 1 else 'signs differ'
                    rows.append([sp, t.replace('h', ' h'), nm]
                                + ['-' if not np.isfinite(v) else '%+.0f%%' % v for v in vals]
                                + ['-' if not np.isfinite(mv) else '%+.0f%%' % mv,
                                   '-' if not np.isfinite(sdv) else '%.0f' % sdv,
                                   str(n), desc])
        grid_table(doc, ['Organism', 'Time', 'Area'] + labels + ['mean', 'SD', 'n', 'signs'],
                   rows, caption='Table 4.%d - anti-adhesion on %s, per batch and pooled. '
                   '"-" = excluded; n = batches contributing; "signs" compares the sign of the '
                   'per-batch values using |AA| > 5 %% as a threshold (arithmetic description of '
                   'the numbers, not a test).' % (num, tag))

    rows = []
    for sp in ('E. coli', 'S. aureus'):
        for ar, nm in (('1', 'LIPSS'), ('2', 'Nanopillar')):
            for t in TIMES:
                ps = [welch(raws[lb].get(sp), t, ar) for lb in labels]
                mv, sdv, n = pooled(sp, t, ar)
                rows.append([sp, t.replace('h', ' h'), nm] + [fmt_p(p) for p in ps]
                            + ['-' if not np.isfinite(mv) else '%+.0f%%' % mv])
    grid_table(doc, ['Organism', 'Time', 'Area'] + ['p %s' % lb for lb in labels] + ['AA mean'],
               rows, caption='Table 4.7 - Welch t-test on total cells, control vs treated, per '
               'batch (uncorrected, n = 5 fields).', fontsize=7.5)

    # ---------------------------------------------------------------- 5. data facts
    doc.add_heading('5. Dataset facts', level=1)
    facts = ['Per batch: 240 images, 24 condition groups of 5 fields each. Groups with fewer than '
             '5 fields after exclusion: %d.' % int(
                 sum(1 for sp in ('E. coli', 'S. aureus') for t in TIMES for ar in AREAS
                     if 0 < (rv(labels[0], sp, t, ar, 'n') if np.isfinite(
                         rv(labels[0], sp, t, ar, 'n')) else 0) < 5))]
    for sp in ('E. coli', 'S. aureus'):
        ctrl = {lb: [rv(lb, sp, t, 'c', 'total_cm2') for t in TIMES] for lb in labels}
        lo = min(v for lab in ctrl.values() for v in lab if np.isfinite(v))
        hi = max(v for lab in ctrl.values() for v in lab if np.isfinite(v))
        ratios = []
        for t in TIMES:
            v = [rv(lb, sp, t, 'c', 'total_cm2') for lb in labels]
            v = [x for x in v if np.isfinite(x) and x > 0]
            if len(v) > 1:
                ratios.append(max(v) / min(v))
        facts.append('%s, control area: total density range %.1f-%.1f x1e5 cells/cm2 across all '
                     'batches and incubation points; largest between-batch ratio for one '
                     'incubation point %.2fx.' % (sp, lo / 1e5, hi / 1e5, max(ratios) if ratios else np.nan))
    add_bullets(doc, facts)
    add_bullets(doc, [
        'Imaging format: batch 1 stored the green and red channels as separate single-channel '
        'images; batches 2 and 3 store both channels in each file (green in channel 1, red in '
        'channel 0, with crosstalk of the order of 1 % of the pixels in the non-target channel). '
        'The counting macro reads only the named channel.',
        'Objects larger than 500 px are excluded by the size filter, so a field containing merged '
        'clusters is counted as the number of objects in the 3-500 px window.',
        'Exposure metadata is not present in the TIFF files, so exposure time per image cannot be '
        'reported.',
    ])

    # ---------------------------------------------------------------- 6. provenance
    doc.add_heading('6. Provenance', level=1)
    add_bullets(doc, [
        'Raw images: Live_Dead Bacteria Test/650nm-20261014/<species>/<batch>/<incubation>/ '
        '(240 images per batch).',
        'Counts: imagej_counts*/<species>/livedead_raw_counts.csv (one table per batch).',
        'Figures, summary CSVs, exclusion registry: E:\\LabToolbox\\output\\650_repeat3_%s\\figs.'
        % a.date,
        'Scripts: recount_650_repeat2_imagej.py (counting), plot_650_multirepeat.py (figures), '
        'build_650_exclusion_registry.py (exclusions), build_650_data_report.py (this report). '
        'Repository github.com/DrFermion/LabToolbox.',
    ])

    docx_path = os.path.join(SCRATCH, '650nm_data_report_%s.docx' % a.date)
    pdf_path = a.out or os.path.join(SCRATCH, '650nm_data_report_%s.pdf' % a.date)
    doc.save(docx_path)
    print('docx saved:', docx_path)
    try:
        import win32com.client
        word = win32com.client.Dispatch('Word.Application')
        word.Visible = False
        try:
            d = word.Documents.Open(os.path.abspath(docx_path))
            d.SaveAs(os.path.abspath(pdf_path), FileFormat=17)
            d.Close(False)
        finally:
            word.Quit()
        print('pdf saved:', pdf_path, os.path.getsize(pdf_path), 'bytes')
        try:
            from pypdf import PdfReader
            rd = PdfReader(pdf_path)
            txt = ''.join((pg.extract_text() or '') for pg in rd.pages)
            print('pages:', len(rd.pages), '| chars:', len(txt), '| has Table 4.7:',
                  'Table 4.7' in txt)
        except Exception as e:
            print('pdf check skipped:', e)
    except Exception as e:
        print('Word COM export failed:', e)
    dst = os.path.join(fig, os.path.basename(docx_path))
    if os.path.abspath(dst) != os.path.abspath(docx_path):
        shutil.copy2(docx_path, dst)
        print('copied ->', dst)


if __name__ == '__main__':
    main()
