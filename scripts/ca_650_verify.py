# -*- coding: utf-8 -*-
"""Verify the updated 650 nm workbook and mirror the outputs."""
import os
import shutil

import openpyxl

BASE = (r"F:\OneDrive - University of Dundee\Svetlana Zolotovskaya (Staff) 的文件 - "
        r"Ruinong Pan_Antibacterial Surfaces\Contact Angle and Surface energy\650")
MIRROR = r"E:\LabToolbox\output\ca_650_20260912"

wb = openpyxl.load_workbook(os.path.join(BASE, "650nm.xlsx"), data_only=True)
for name in ("summary", "raw note", "area1"):
    ws = wb[name]
    print(f"=== {name} ===")
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row):
        vals = [("" if c.value is None else str(c.value)) for c in r[:8]]
        while vals and vals[-1] == "":
            vals.pop()
        if vals:
            print(f"  r{r[0].row:<3}| " + " | ".join(vals))
    print()

os.makedirs(MIRROR, exist_ok=True)
for f in ("650nm.xlsx", "650nm_water_contact_angle.png",
          "650nm_water_contact_angle_first_repeat.png",
          "650nm_water_contact_angle_repeat_drift.png"):
    shutil.copy2(os.path.join(BASE, f), os.path.join(MIRROR, f))
print("镜像到:", MIRROR)
print("内容:", sorted(os.listdir(MIRROR)))
