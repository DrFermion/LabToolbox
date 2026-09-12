# -*- coding: utf-8 -*-
"""650 nm water contact angle — keep the workbook and the plots in sync.

Workflow (agreed 2026-09-12): 主人直接在 xlsx 的 area 表里手填每天的数据 —— 每张表底部预留了
20 行"待填区"(淡黄底 + 边框)。荧荧只读文件重画图, **不再需要辨认手写照片**。

每次跑这个脚本都会:
  1. 从四张 area 表读出所有 day (Day 0/1/2/...) —— 天数是动态的, 以后加 Day 3、Day 4 自动进图;
  2. 重算每组的 mean ± SD, 写回 summary 的数值列 (日期/温度/RH 保留主人在表里写的);
  3. 从 summary 读温度/RH/日期, 重画三张图:
       650nm_water_contact_angle.png              mean ± SD (全部重复)
       650nm_water_contact_angle_first_repeat.png 每组只取第 1 次重复 (干燥不完全的参考)
       650nm_water_contact_angle_repeat_drift.png 单次会话内 repeat 1→2→3 的漂移
  4. 确保每张 area 表底部有 20 行空白待填区。

一次性动作: Day 2 (12/9) 的手写转录值已在表里, 脚本检测到就跳过, 不会重复写。
"""
import os
import re
import statistics as st

import openpyxl
from openpyxl.styles import Border, PatternFill, Side

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

BASE = (r"F:\OneDrive - University of Dundee\Svetlana Zolotovskaya (Staff) 的文件 - "
        r"Ruinong Pan_Antibacterial Surfaces\Contact Angle and Surface energy\650")
XLSX = os.path.join(BASE, "650nm.xlsx")

GROUPS = ["area1", "area1-R90", "area2", "control"]
LABEL = {"area1": "area1", "area1-R90": "area1 (rot 90°)", "area2": "area2", "control": "control"}
COLOR = {"area1": "#1565C0", "area1-R90": "#EF6C00", "area2": "#6A1B9A", "control": "#546E7A"}
MARK = {"area1": "o", "area1-R90": "s", "area2": "^", "control": "D"}

ENTRY_ROWS = 20           # 底部预留的空白待填行数
ENTRY_FILL = "FFFDE7"     # 淡黄 = 待填

# 一次性: Day 2 (12/9, 23 °C, 44% RH) 的手写转录, 四列 = area1 | area1(rot 90°) | area2 | control
DAY2 = {
    "area1":     [(82.1, 82.5), (67.6, 69.6), (57.7, 60.3)],
    "area1-R90": [(65.6, 67.3), (58.1, 55.8), (52.3, 51.7)],
    "area2":     [(64.4, 64.7), (54.0, 54.4), (51.1, 51.9)],
    "control":   [(65.1, 66.0), (67.1, 68.3), (65.8, 64.2)],
}


def avg(pair):
    vals = [v for v in pair if v is not None]
    return sum(vals) / len(vals)


def day_key(d):
    m = re.search(r"(\d+)", str(d))
    return int(m.group(1)) if m else 999


def is_day(v):
    return isinstance(v, str) and v.strip().lower().startswith("day")


def last_used_row(ws):
    last = 0
    for r in range(3, ws.max_row + 1):
        for c in (1, 2, 3, 4, 5):
            v = ws.cell(r, c).value
            if v is not None and not (isinstance(v, str) and v.strip() == ""):
                last = r
    return last


def read_sheet(ws):
    """{day: {repeat: value}} from one area sheet."""
    out = {}
    for r in range(3, ws.max_row + 1):
        day, rep = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not is_day(day) or rep is None:
            continue
        if isinstance(rep, str) and rep.strip().lower().startswith("average"):
            continue
        try:
            rep_i = int(rep)
        except (TypeError, ValueError):
            continue
        nums = [float(ws.cell(r, c).value) for c in (3, 4)
                if isinstance(ws.cell(r, c).value, (int, float))]
        if nums:
            out.setdefault(str(day).strip(), {})[rep_i] = sum(nums) / len(nums)
    return out


def stat(data, day):
    vals = list(data.get(day, {}).values())
    if not vals:
        return None, None
    return st.mean(vals), (st.stdev(vals) if len(vals) > 1 else 0.0)


def read_conditions(wb):
    """{day: (date, temp, rh)} from the summary sheet (as typed by the user)."""
    ws = wb["summary"]
    out = {}
    for r in range(3, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if is_day(d):
            date = ws.cell(r, 2).value
            if hasattr(date, "strftime"):
                date = f"{date.day}/{date.month}"
            out[str(d).strip()] = (date, ws.cell(r, 3).value, ws.cell(r, 4).value)
    return out


def ensure_entry_rows(ws, n=ENTRY_ROWS):
    """Keep n blank, bordered, pale-yellow rows directly under the data (rolling buffer).

    Rows are painted only after the last row that holds a value, so typed data is never
    touched; if the user fills some of the block, the next run tops it back up to n.
    """
    first = last_used_row(ws) + 1
    thin = Side(style="thin", color="B0BEC5")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor=ENTRY_FILL)
    blank = 0
    for i in range(n):
        for c in range(1, 6):
            cell = ws.cell(first + i, c)
            if cell.value is None:
                cell.border = bd
                cell.fill = fill
                blank += 1
    return first, blank


def append_day2(ws, pairs):
    start = last_used_row(ws) + 2
    for i, (left, right) in enumerate(pairs):
        r = start + i
        ws.cell(r, 1, "Day 2")
        ws.cell(r, 2, i + 1)
        ws.cell(r, 3, left)
        ws.cell(r, 4, right)
        ws.cell(r, 5, round(avg((left, right)), 2))
    ravg = start + len(pairs)
    ws.cell(ravg, 1, "Day 2")
    ws.cell(ravg, 2, "average")
    ws.cell(ravg, 5, round(st.mean(avg(p) for p in pairs), 2))
    return start, ravg


def draw(path, values, days, cond, title, note, ymax=None):
    fig, ax = plt.subplots(figsize=(7.8, 5.4))
    xs = list(range(len(days)))
    for g in GROUPS:
        ys = [values[g].get(d, (None, 0))[0] for d in days]
        es = [values[g].get(d, (None, 0))[1] for d in days]
        ax.errorbar(xs, ys, yerr=es, color=COLOR[g], marker=MARK[g], ls="-", lw=1.6,
                    ms=6, capsize=4, label=LABEL[g])
        for xx, yy in zip(xs, ys):
            if yy is not None:
                ax.annotate(f"{yy:.1f}", (xx, yy), xytext=(6, 7),
                            textcoords="offset points", fontsize=8, color=COLOR[g])
    ax.set_xticks(xs)
    ax.set_xticklabels([f"{d} ({cond.get(d, ('',))[0]})" if cond.get(d, ('',))[0] else d
                        for d in days])
    for i, d in enumerate(days):
        c = cond.get(d)
        if c and c[1] is not None and c[2] is not None:
            ax.annotate(f"{c[1]} °C, {c[2]}% RH", (i, 0), xytext=(0, -34),
                        textcoords="offset points", ha="center", fontsize=8, color="#555555")
    ax.set_ylabel("Water contact angle (deg)")
    ax.set_ylim(0, ymax or 100)
    ax.set_title(title, fontsize=11)
    ax.grid(axis="y", alpha=0.25)
    ax.legend(title="Surface area", fontsize=9, title_fontsize=9, loc="upper left")
    if note:
        ax.text(0.99, 0.02, note, transform=ax.transAxes, ha="right", va="bottom",
                fontsize=7.5, color="#666666", style="italic")
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)
    return path


def main():
    wb = openpyxl.load_workbook(XLSX)
    print("工作表:", wb.sheetnames)

    # 1. one-off: Day 2 handwriting, only if absent
    for g in GROUPS:
        ws = wb[g]
        have = {str(ws.cell(r, 1).value).strip() for r in range(3, ws.max_row + 1)}
        if "Day 2" in have:
            continue
        start, ravg = append_day2(ws, DAY2[g])
        print(f"  {g}: Day 2 写入 r{start}-r{ravg}")

    # 2. read every day present (dynamic)
    data = {g: read_sheet(wb[g]) for g in GROUPS}
    all_days = sorted({d for g in GROUPS for d in data[g]}, key=day_key)
    cond = read_conditions(wb)
    print("天数:", all_days, " 条件:", {d: cond.get(d) for d in all_days})

    # 3. summary stays in sync: add a row for any day that only exists in the area sheets,
    #    then rewrite the numeric columns and the Δ block below the day rows
    ws = wb["summary"]

    def day_rows():
        return sorted([(r, str(ws.cell(r, 1).value).strip()) for r in range(3, ws.max_row + 1)
                       if is_day(ws.cell(r, 1).value)], key=lambda t: t[0])

    rows = day_rows()
    anchor = rows[-1][0] if rows else 5
    have = {d for _, d in rows}
    for d in all_days:
        if d in have:
            continue
        anchor += 1
        ws.insert_rows(anchor)
        ws.cell(anchor, 1, d)
        have.add(d)
    rows = day_rows()
    for r, d in rows:
        for j, g in enumerate(GROUPS):
            m, s = stat(data[g], d)
            if m is not None:
                ws.cell(r, 5 + j, f"{m:.2f} ± {s:.2f}")
    # clear whatever sits under the day rows (old Δ block) and rewrite it
    for r in range(rows[-1][0] + 1, ws.max_row + 1):
        for c in range(1, 9):
            ws.cell(r, c).value = None
    r = rows[-1][0] + 1
    for title, (d1, d2) in {"Δ (Day 1 − Day 0)": ("Day 0", "Day 1"),
                            "Δ (Day 2 − Day 1)": ("Day 1", "Day 2"),
                            "Δ (Day 2 − Day 0)": ("Day 0", "Day 2")}.items():
        if d1 not in all_days or d2 not in all_days:
            continue
        ws.cell(r, 1, title)
        for j, g in enumerate(GROUPS):
            m1, _ = stat(data[g], d1)
            m2, _ = stat(data[g], d2)
            ws.cell(r, 5 + j, f"{m2 - m1:+.2f}")
        r += 1

    # 4. 20 blank entry rows at the bottom of every area sheet
    for g in GROUPS:
        first, blank = ensure_entry_rows(wb[g])
        print(f"  {g}: 待填区 r{first}-r{first + ENTRY_ROWS - 1} (画框空白单元格 {blank} 个)")

    wb.save(XLSX)
    print("已写回:", XLSX)

    # 5. figures
    mean_sd = {g: {d: stat(data[g], d) for d in all_days} for g in GROUPS}
    draw(os.path.join(BASE, "650nm_water_contact_angle.png"), mean_sd, all_days, cond,
         "650 nm sample — water contact angle (mean ± SD, n = 3)", "")
    first_rep = {g: {d: (data[g][d].get(1), 0.0) for d in all_days if 1 in data[g].get(d, {})}
                 for g in GROUPS}
    draw(os.path.join(BASE, "650nm_water_contact_angle_first_repeat.png"), first_rep, all_days, cond,
         "650 nm sample — water contact angle, 1st repeat only (mean of left/right)",
         "first drop of each measurement session")

    fig, ax = plt.subplots(figsize=(7.8, 5.4))
    for g in GROUPS:
        for i, d in enumerate(all_days):
            ys = [data[g][d].get(k) for k in (1, 2, 3)]
            if any(v is None for v in ys):
                continue
            ax.plot([k + i * 3.6 for k in range(3)], ys, color=COLOR[g], marker=MARK[g],
                    ls="-" if g != "control" else "--", lw=1.4, ms=5, alpha=0.85,
                    label=LABEL[g] if i == 0 else None)
    ax.set_xticks([i * 3.6 + 1 for i in range(len(all_days))])
    ax.set_xticklabels(all_days)
    ax.set_xlabel("repeat 1 → 2 → 3 within each session")
    ax.set_ylabel("Water contact angle (deg)")
    ax.set_title("650 nm — repeat-to-repeat drift within a session", fontsize=11)
    ax.grid(axis="y", alpha=0.25)
    ax.legend(title="Surface area", fontsize=9, title_fontsize=9, loc="upper left")
    fig.tight_layout()
    drift = os.path.join(BASE, "650nm_water_contact_angle_repeat_drift.png")
    fig.savefig(drift, dpi=200)
    plt.close(fig)

    print()
    print("mean ± SD (all repeats) / repeat 1 only:")
    for g in GROUPS:
        row = "  %-16s" % LABEL[g]
        for d in all_days:
            m, s = mean_sd[g][d]
            v = data[g].get(d, {}).get(1)
            row += f"  {d}: {m:5.1f}±{s:4.1f}" if m is not None else f"  {d}:   -  "
            row += f" [1st {v:5.1f}]" if v is not None else ""
        print(row)
    print()
    for p in ("650nm_water_contact_angle.png", "650nm_water_contact_angle_first_repeat.png",
              "650nm_water_contact_angle_repeat_drift.png"):
        print("图:", os.path.join(BASE, p))


if __name__ == "__main__":
    main()
