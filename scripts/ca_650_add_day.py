# -*- coding: utf-8 -*-
"""Add one day of 650 nm contact-angle readings to the workbook (generic).

Usage:
    python ca_650_add_day.py --day "Day 6" --date "16/9" --temp 24 --rh 39 \
        --json C:/path/day6.json

JSON format (left/right readings per repeat, in workbook column order):
    {"area1":     [[98.3, 98.4], [72.4, 74.8], [72.7, 75.3]],
     "area1-R90": [[82.2, 83.3], [71.5, 72.7], [66.2, 65.3]],
     "area2":     [[79.2, 79.3], [60.2, 62.4], [54.6, 53.7]],
     "control":   [[69.0, 70.4], [65.5, 65.6], [61.2, 62.6]]}

Writes a Day block (3 repeats + average) into each area sheet, then makes sure the
summary sheet has a row for that day carrying the date / temperature / RH. The numeric
columns are filled by ca_650_update.py (run it afterwards).
"""
import argparse
import json
import os
import re
import shutil
import statistics as st
import time

import openpyxl

BASE = (r"F:\OneDrive - University of Dundee\Svetlana Zolotovskaya (Staff) 的文件 - "
        r"Ruinong Pan_Antibacterial Surfaces\Contact Angle and Surface energy\650")
XLSX = os.path.join(BASE, "650nm.xlsx")
GROUPS = ["area1", "area1-R90", "area2", "control"]


def is_day(v):
    return isinstance(v, str) and v.strip().lower().startswith("day")


def last_used_row(ws):
    last = 0
    for r in range(3, ws.max_row + 1):
        for c in (1, 2, 3, 4, 5):
            v = ws.cell(r, c).value
            if v is not None and not (isinstance(v, str) and not v.strip()):
                last = r
    return last


def avg(pair):
    vals = [float(v) for v in pair if v is not None]
    return sum(vals) / len(vals)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--day", required=True, help='label, e.g. "Day 6"')
    ap.add_argument("--date", required=True, help='e.g. "16/9"')
    ap.add_argument("--temp", type=float, default=None)
    ap.add_argument("--rh", type=float, default=None)
    ap.add_argument("--json", required=True, help="path to the readings JSON")
    args = ap.parse_args()

    data = json.load(open(args.json, encoding="utf-8"))
    day = args.day.strip()
    for g in GROUPS:
        if len(data.get(g, [])) != 3:
            raise SystemExit(f"{g}: 需要 3 次重复, 得到 {len(data.get(g, []))}")

    bak = XLSX + f".bak_before_{day.replace(' ', '')}_{time.strftime('%Y%m%d-%H%M')}"
    shutil.copy2(XLSX, bak)
    print("备份:", os.path.basename(bak))

    wb = openpyxl.load_workbook(XLSX)

    # 1) day blocks in every area sheet
    for g in GROUPS:
        ws = wb[g]
        have = {str(ws.cell(r, 1).value).strip() for r in range(3, ws.max_row + 1)}
        if day in have:
            print(f"  {g}: {day} 已存在, 跳过")
            continue
        start = last_used_row(ws) + 1
        for i, pair in enumerate(data[g]):
            r = start + i
            ws.cell(r, 1, day)
            ws.cell(r, 2, i + 1)
            ws.cell(r, 3, float(pair[0]))
            ws.cell(r, 4, float(pair[1]) if pair[1] is not None else None)
            ws.cell(r, 5, round(avg(pair), 2))
        ravg = start + len(data[g])
        ws.cell(ravg, 1, day)
        ws.cell(ravg, 2, "average")
        ws.cell(ravg, 5, round(st.mean(avg(p) for p in data[g]), 2))
        print(f"  {g}: 写入 r{start}-r{ravg}")

    # 2) summary: make sure the day has a row with its conditions
    ws = wb["summary"]
    rows = [(r, str(ws.cell(r, 1).value).strip()) for r in range(3, ws.max_row + 1)
            if is_day(ws.cell(r, 1).value)]
    labels = [d for _, d in rows]
    if day not in labels:
        # 按天数插到正确位置 (不能简单追加到末尾, 否则 Day 3 会排到 Day 8 后面)
        num = int(re.search(r"(\d+)", day).group(1))
        earlier = [rr for rr, d in rows
                   if re.search(r"(\d+)", d) and int(re.search(r"(\d+)", d).group(1)) < num]
        anchor = max(earlier) if earlier else rows[0][0] - 1
        ws.insert_rows(anchor + 1)
        r = anchor + 1
        ws.cell(r, 1, day)
        print(f"  summary: 新增 {day} 行 (r{r}, 按天数插入)")
    else:
        r = dict((d, rr) for rr, d in rows)[day]
    ws.cell(r, 2, args.date)
    if args.temp is not None:
        ws.cell(r, 3, args.temp)
    if args.rh is not None:
        ws.cell(r, 4, args.rh)
    print(f"  summary r{r}: {day} | {args.date} | {args.temp} | {args.rh}")

    wb.save(XLSX)
    print("已写回:", XLSX)


if __name__ == "__main__":
    main()
