# -*- coding: utf-8 -*-
"""
recount_livedead_515_imagej.py — 515nm 两组 LIVE/DEAD 重计数 (ImageJ 引擎) + 统计报告
- 515nm (盖玻片): 自动忽略数据不全的 repeat2 (10x 图), repeat3 空目录自动跳过
- 515nm-filter_paper (滤纸): 全量 (repeat1)
- 引擎: 经典 ImageJ 1.54 headless (自适应阈值 背景众数+40, Analyze Particles 3-500px)
- 40X 视野面积: 46946.10 µm² (密度换算 cells/cm²)
输出: scripts/output/recount_livedead_515/<组>/... + 汇总 docx 报告
"""
import os
import sys
import datetime
import numpy as np
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from labtoolbox.livedead_cellcounter.cellcounter import LiveDeadCellCounter, ImageJEngine

BASE = (r"F:\OneDrive - University of Dundee\Svetlana Zolotovskaya (Staff) 的文件 - Ruinong Pan_"
        r"Antibacterial Surfaces\Live_Dead Bacteria Test")
GROUPS = [
    {"name": "515nm_coverslip", "root": os.path.join(BASE, "515nm"),
     "skip": ("repeat2",), "area_um2": 46946.10},
    {"name": "515nm_filter_paper", "root": os.path.join(BASE, "515nm-filter_paper"),
     "skip": (), "area_um2": 46946.10},
]
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output",
                   "recount_livedead_515")


def density_per_cm2(count, area_um2):
    """count per image -> cells/cm² (1 cm² = 1e8 µm²)"""
    if not area_um2 or pd.isna(count):
        return np.nan
    return count / area_um2 * 1e8


def main():
    os.makedirs(OUT, exist_ok=True)
    engine = ImageJEngine()
    group_reports = {}
    for g in GROUPS:
        print("\n" + "=" * 60)
        print("组: %s   (跳过: %s)" % (g["name"], g["skip"] or "无"))
        od = os.path.join(OUT, g["name"])
        analyzer = LiveDeadCellCounter(engine=engine, area_um2=g["area_um2"])
        rep = analyzer.report(g["root"], output_dir=od, skip_repeats=g["skip"])
        df = rep["raw"].copy()
        # 密度列 (每视野面积换算)
        df["live_per_cm2"] = df["live"].apply(density_per_cm2, area_um2=g["area_um2"])
        df["dead_per_cm2"] = df["dead"].apply(density_per_cm2, area_um2=g["area_um2"])
        df["total_per_cm2"] = df["total"].apply(density_per_cm2, area_um2=g["area_um2"])
        rep["raw"] = df
        # 保存带密度的 raw
        df.to_csv(os.path.join(od, "livedead_raw_counts.csv"),
                  index=False, encoding="utf-8-sig")
        summ = rep["summary"].copy()
        summ["live_cm2"] = summ["live_mean"].apply(density_per_cm2, area_um2=g["area_um2"])
        summ["total_cm2"] = summ["total_mean"].apply(density_per_cm2, area_um2=g["area_um2"])
        summ = summ[["area_label", "time", "n", "live_mean", "live_sd", "dead_mean",
                     "dead_sd", "total_mean", "total_sd", "viability_pct",
                     "live_cm2", "total_cm2"]]
        summ.to_csv(os.path.join(od, "livedead_summary.csv"),
                    index=False, encoding="utf-8-sig")
        rep["summary"] = summ
        rep["summ_csv"] = os.path.join(od, "livedead_summary.csv")
        group_reports[g["name"]] = {"rep": rep, "cfg": g}
        print("汇总 (mean ± SD per image):")
        pd.set_option("display.width", 200)
        print(summ.round(1).to_string(index=False))

    # ---- docx 报告 ----
    try:
        from docx import Document
        from docx.shared import Pt, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        print("python-docx 缺失, 跳过 docx 报告")
        return

    doc = Document()
    doc.add_heading("LIVE/DEAD Bacteria Re-count Report (ImageJ)", 0)
    doc.add_paragraph("Date: %s" % datetime.date.today().isoformat())
    doc.add_heading("Methods", level=1)
    doc.add_paragraph(
        "Counting engine: ImageJ 1.54 (classic, headless batch) with adaptive threshold "
        "(background histogram mode + 40) and Analyze Particles (size 3-500 px). "
        "Green (g) images counted for live bacteria, red (r) images for dead; "
        "g/r pairs of the same field merged per field. "
        "Field area (40x): 46946.10 um2. "
        "515nm coverslip group: repeat2 excluded (incomplete 10x data); empty repeat3 skipped. "
        "515nm filter-paper group: repeat1 analysed.")

    for name, gr in group_reports.items():
        cfg, rep = gr["cfg"], gr["rep"]
        doc.add_heading("Group: %s" % name, level=1)
        doc.add_paragraph("Data root: %s" % cfg["root"])
        doc.add_paragraph("Repeat included: %s" % ("repeat1" if not cfg["skip"] else
                                                   "repeat1 (repeat2 excluded)"))
        # summary table
        doc.add_heading("Summary (counts per field, mean +/- SD)", level=2)
        summ = rep["summary"]
        t = doc.add_table(rows=1, cols=len(summ.columns))
        t.style = "Light Grid Accent 1"
        for j, c in enumerate(summ.columns):
            t.rows[0].cells[j].text = str(c)
        for _, r in summ.iterrows():
            cells = t.add_row().cells
            for j, c in enumerate(summ.columns):
                v = r[c]
                cells[j].text = ("%.1f" % v) if isinstance(v, float) else str(v)
        # figure
        if os.path.exists(rep["figure"]):
            doc.add_picture(rep["figure"], width=Inches(6.0))
        # anova
        aov = rep["anova"]
        if aov is not None and len(aov):
            doc.add_heading("ANOVA (log10 live counts, area x time)", level=2)
            t2 = doc.add_table(rows=1, cols=len(aov.columns))
            t2.style = "Light Grid Accent 1"
            for j, c in enumerate(aov.columns):
                t2.rows[0].cells[j].text = str(c)
            for _, r in aov.iterrows():
                cells = t2.add_row().cells
                for j, c in enumerate(aov.columns):
                    cells[j].text = "%.3g" % r[c] if isinstance(r[c], float) else str(r[c])
        # output files
        doc.add_heading("Output files", level=2)
        for k in ("raw_csv", "summ_csv", "anova_csv", "figure"):
            doc.add_paragraph("%s: %s" % (k, rep.get(k)))

    report_path = os.path.join(OUT, "LIVE_DEAD_515_recount_report.docx")
    doc.save(report_path)
    print("\n报告: %s" % report_path)


if __name__ == "__main__":
    main()
