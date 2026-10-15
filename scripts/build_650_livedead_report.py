# -*- coding: utf-8 -*-
"""build_650_livedead_report.py — 650 nm LIVE/DEAD repeat-1 汇总 PDF 报告。

数据源: <共享夹>/.../650nm-20261014/imagej_counts/ 的汇总 CSV + 图
流程  : 读数 -> python-docx 生成报告 -> Word COM 导出 PDF -> 分发(共享夹/个人夹/镜像)
用法  : C:/Users/PC/miniconda3/python.exe build_650_livedead_report.py
"""
import os
import shutil
import subprocess
import pandas as pd
from scipy import stats
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

DS = ('F:/OneDrive - University of Dundee/Svetlana Zolotovskaya (Staff) 的文件 - '
      'Ruinong Pan_Antibacterial Surfaces/Live_Dead Bacteria Test/650nm-20261014')
IC = os.path.join(DS, 'imagej_counts')
PERSONAL = ('F:/OneDrive - University of Dundee/Ruinong_Pan_Personal/'
            '650nm_LIVE-DEAD_repeat1_20261015')
MIRROR = 'E:/LabToolbox/output/livedead_650nm_20261014'
SCRATCH = 'C:/Users/PC/AppData/Local/hermes/cache/scratch'
FIG = os.path.join(IC, '650nm_livedead_timecourse_repeat1.png')
PREVIEW = os.path.join(IC, 'S.aureus', 'qc', 'saureus5h_channel_preview.png')
SUMMARY = os.path.join(IC, '650nm_livedead_timecourse_summary.csv')
EXPOSURE = os.path.join(IC, 'exposure_log.csv')
PDF_NAME = '650nm_LIVE-DEAD_repeat1_summary_20261015.pdf'

TIMES = ['3h', '5h', '7h', '24h']
REGS = ['c', '1', '2']
REG_LAB = {'c': 'Control', '1': 'LIPSS', '2': 'Nanopillar'}
SPECIES = [('E. coli', 'E.coli'), ('S. aureus', 'S.aureus')]

S = pd.read_csv(SUMMARY)
EXP = pd.read_csv(EXPOSURE)
KS = 1e8 / 46946.10  # counts per field -> cells/cm2


def val(sp, t, reg, col):
    g = S.query('species == @sp and time == @t and region == @reg')
    return float(g[col].iloc[0]) if len(g) else None


def anova_p(sp_dir, t):
    raw = pd.read_csv(os.path.join(IC, sp_dir, 'livedead_raw_counts.csv'))
    raw = raw[raw.time.astype(str) == t].dropna(subset=['live', 'dead'])
    grp = [raw[raw.area.astype(str) == reg] for reg in REGS]
    if any(len(g) < 2 for g in grp):
        return None
    viab = [g['live'] / (g['live'] + g['dead']) * 100 for g in grp]
    F, p = stats.f_oneway(*viab)
    return p


def fmt_p(p):
    if p is None:
        return '—'
    if p < 0.001:
        return '< 0.001'
    return '%.3f' % p


def exposure_exceptions():
    notes = []
    for (sp, t), g in EXP.groupby(['species', 'time']):
        vals = g['exposure_ms'].astype(int)
        mode = int(vals.mode().iloc[0])
        for area, ga in g.groupby('area'):
            vv = ga['exposure_ms'].astype(int)
            dev = sorted(set(vv[vv != mode]))
            if dev:
                frac = float(vv.isin(dev).mean())
                notes.append('%s %s %s %s ms%s' % (sp, t, area, '/'.join(map(str, dev)),
                                                   '' if frac == 1 else ' (part)'))
    return notes


# ---------------- docx ----------------
doc = Document()
for sec in doc.sections:
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
st = doc.styles['Normal']
st.font.name = 'Arial'
st.font.size = Pt(10)

h = doc.add_heading('650 nm textured stainless steel - bacterial LIVE/DEAD summary (repeat 1)', level=0)
p = doc.add_paragraph('Prepared 15 October 2026  ·  Dataset: Live_Dead Bacteria Test/650nm-20261014  ·  '
                      'E. coli and S. aureus  ·  timepoints 3 / 5 / 7 / 24 h  ·  n = 5 fields per condition')
p.alignment = WD_ALIGN_PARAGRAPH.LEFT

doc.add_heading('1. Scope and method', level=1)
for line in [
    'Sample: 650 nm laser-textured stainless-steel coupon with three areas - Control (c), LIPSS (1), '
    'Nanopillar (2). Two organisms were incubated on the surfaces for 3, 5, 7 and 24 h and imaged after '
    'LIVE/DEAD staining (green = live, red = dead) at 40x magnification (field area 46,946.10 um2; '
    '1392 x 1040 px). Five fields per area and timepoint were recorded (240 images in total).',
    'Counting: ImageJ 1.54 "Analyze Particles" with an adaptive threshold (background mode + 40, particle '
    'size filter 3-500 px); green and red images were paired per field. An independent re-implementation of '
    'the same detection geometry reproduced every image count (0/240 mismatches; 10,841 objects for '
    'E. coli, 139,862 for S. aureus).',
    'Derived values: density = count / 46,946.10 um2 x 1e8 (cells/cm2); viability = pooled live / '
    '(live + dead) per condition. One-way ANOVA across areas was run per organism and timepoint.']:
    doc.add_paragraph(line)

doc.add_heading('2. Results', level=1)
doc.add_paragraph('Figure 1. LIVE/DEAD summary - stacked bars (green = live, red = dead), '
                  'density in cells/cm2; upper panels E. coli, lower panels S. aureus. '
                  'Bar labels = pooled viability. "?" marks a QC-flagged point (see Section 3).')
doc.add_picture(FIG, width=Cm(16.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

for sp, sp_dir in SPECIES:
    doc.add_heading('Viability (%%) - %s' % sp, level=2)
    tb = doc.add_table(rows=1, cols=5)
    tb.style = 'Light Grid Accent 1'
    for i, htxt in enumerate(['Time', 'Control', 'LIPSS', 'Nanopillar', 'ANOVA p (areas)']):
        tb.rows[0].cells[i].text = htxt
    for t in TIMES:
        c = tb.add_row().cells
        c[0].text = t.replace('h', ' h')
        for j, reg in enumerate(REGS):
            v = val(sp, t, reg, 'viability_pooled')
            c[j + 1].text = '—' if v is None else ('%.1f' % v if not (sp == 'S. aureus' and t == '5h' and reg == '2') else '%.1f ?' % v)
        c[4].text = fmt_p(anova_p(sp_dir, t))
    doc.add_paragraph()
    doc.add_heading('Live-cell density (x 10^5 cells/cm2) - %s' % sp, level=2)
    tb = doc.add_table(rows=1, cols=5)
    tb.style = 'Light Grid Accent 1'
    for i, htxt in enumerate(['Time', 'Control', 'LIPSS', 'Nanopillar', 'ANOVA p (viability)']):
        tb.rows[0].cells[i].text = htxt
    for t in TIMES:
        c = tb.add_row().cells
        c[0].text = t.replace('h', ' h')
        for j, reg in enumerate(REGS):
            v = val(sp, t, reg, 'live_cm2')
            c[j + 1].text = '—' if v is None else ('%.2f' % (v / 1e5))
        c[4].text = fmt_p(anova_p(sp_dir, t))
    doc.add_paragraph()

doc.add_paragraph('Highlights:', style='Heading 2')
for line in [
    'E. coli viability (Control / LIPSS / Nanopillar): 61/53/59 -> 54/27/59 -> 43/38/32 -> 84/74/83 %.',
    'S. aureus viability: 93/73/57 -> 77/67/flagged -> 74/61/84 -> 75/64/82 %.',
    'At 24 h, LIPSS showed the lowest viability in both organisms; between-area differences were '
    'significant at most timepoints (p <= 0.01), except E. coli at 3 h (p = 0.054) and 7 h (ns).']:
    doc.add_paragraph(line, style='List Bullet')

doc.add_heading('3. Notes, QC and limitations', level=1)
for line in [
    'Single repeat. Only repeat 1 of three planned repeats is available; repeats 2-3 are pending. No '
    'cross-repeat statistics yet - treat time trends as provisional.',
    'Exposure time varies between and within timepoints (read from the TIFF metadata, see '
    'exposure_log.csv): all E. coli 24 h images were taken at 700 ms (vs 250-300 ms at earlier '
    'timepoints); S. aureus 24 h LIPSS at 500 ms vs 200 ms for control/nanopillar; S. aureus 5 h '
    'nanopillar at 500 ms. Absolute counts are therefore not directly comparable across sessions; '
    'viability ratios are less affected (both channels share the exposure when paired).',
    'S. aureus 5 h, Nanopillar (marked "?"): the green channel was nearly empty even at twice the '
    'nominal exposure, while the 7 h group imaged at 0.25 s shows abundant green signal. This pattern is '
    'most consistent with an imaging/acquisition problem; the 0.4 % value should not be quoted until '
    'the images are re-checked or re-acquired. Preview of the field: left = 2-g3 (green channel, no '
    'signal), middle = 2-r3 (red channel, dense), right = control c-g1 (green, normal).',
    'Counting excludes clusters larger than 500 px, so counts are slight lower bounds in dense fields '
    '(merged-clump pixels ~4-5 % of signal at 24 h; higher in isolated fields).',
    'E. coli control viability is non-monotonic over time (61 -> 54 -> 43 -> 84 %), suggesting '
    'per-session measurement (staining/imaging) effects; interpret time comparisons together with '
    'repeats 2-3.',
    'One image (S. aureus 5 h, 2-g3) contained no measurable green signal (count = 1 via fallback); '
    'including or excluding it does not change any conclusion.']:
    doc.add_paragraph(line, style='List Bullet')
doc.add_picture(PREVIEW, width=Cm(14.0))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_heading('4. Data and provenance', level=1)
for line in [
    'Dataset: %s (240 TIFF).' % DS,
    'Counts / summaries: imagej_counts/<species>/livedead_raw_counts.csv; combined summary: '
    'imagej_counts/650nm_livedead_timecourse_summary.csv; exposure log: imagej_counts/exposure_log.csv; '
    'QC overlays: imagej_counts/<species>/qc/<time>/ (30 per group).',
    'Scripts: E:\\LabToolbox\\scripts\\plot_650_livedead_timecourse.py (figure), '
    'build_650_livedead_report.py (this report). Repository: github.com/DrFermion/LabToolbox.']:
    doc.add_paragraph(line, style='List Bullet')

docx_path = os.path.join(SCRATCH, '650nm_report.docx')
pdf_path = os.path.join(SCRATCH, PDF_NAME)
doc.save(docx_path)
print('docx saved:', docx_path)

# ------------- Word -> PDF -------------
import win32com.client
word = win32com.client.Dispatch('Word.Application')
word.Visible = False
try:
    d = word.Documents.Open(docx_path)
    d.SaveAs(pdf_path, FileFormat=17)  # wdFormatPDF
    d.Close(False)
finally:
    word.Quit()
print('pdf saved:', pdf_path, os.path.getsize(pdf_path), 'bytes')

# verify
PdfReader = None
for _mod in ('pypdf', 'PyPDF2'):
    try:
        PdfReader = __import__(_mod).PdfReader
        break
    except Exception:
        pass
if PdfReader is None:
    print('WARN: pypdf/PyPDF2 not available; skip PDF text verification')
else:
    rd = PdfReader(pdf_path)
    print('pages:', len(rd.pages))
    txt = ''.join((pg.extract_text() or '') for pg in rd.pages)
    print('text chars:', len(txt), '| contains "Nanopillar":', 'Nanopillar' in txt,
          '| contains "flagged":', 'flagged' in txt)

# ------------- distribute -------------
os.makedirs(PERSONAL, exist_ok=True)
targets = [
    (pdf_path, os.path.join(DS, PDF_NAME)),
    (pdf_path, os.path.join(PERSONAL, PDF_NAME)),
    (pdf_path, os.path.join(MIRROR, PDF_NAME)),
    (FIG, os.path.join(PERSONAL, os.path.basename(FIG))),
    (SUMMARY, os.path.join(PERSONAL, os.path.basename(SUMMARY))),
    (EXPOSURE, os.path.join(PERSONAL, os.path.basename(EXPOSURE))),
    (FIG, os.path.join(MIRROR, os.path.basename(FIG))),
]
for src, dst in targets:
    shutil.copy2(src, dst)
    print('copied ->', dst, os.path.getsize(dst))
print('exposure exceptions:', '; '.join(exposure_exceptions()))
