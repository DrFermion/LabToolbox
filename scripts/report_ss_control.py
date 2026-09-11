# -*- coding: utf-8 -*-
"""
report_ss_control.py — ss-control 单绿染总菌计数 docx 报告 (ImageJ 引擎)

读取 imagej_counts/ss_control_raw.csv + summary CSV, 生成:
  imagej_counts/ss_control_report.docx
含: 方法学 / 逐 repeat 汇总表 / 合并汇总表 / 统计检验表 / 两张图 / 文件清单 / 备注
"""
import os
import datetime
import numpy as np
import pandas as pd
from scipy import stats

BASE = (r"F:\OneDrive - University of Dundee\Svetlana Zolotovskaya (Staff) 的文件 - Ruinong Pan_"
        r"Antibacterial Surfaces\Live_Dead Bacteria Test\ss-control")
OUT = os.path.join(BASE, "imagej_counts")
REPEATS = ("repeat 1", "repeat 2", "repeat 3")
TIMES = ("3h", "5h")

raw = pd.read_csv(os.path.join(OUT, "ss_control_raw.csv"), encoding="utf-8-sig")
by_rep = pd.read_csv(os.path.join(OUT, "ss_control_summary_by_repeat.csv"), encoding="utf-8-sig")
summ = pd.read_csv(os.path.join(OUT, "ss_control_summary.csv"), encoding="utf-8-sig")
clean = raw[~raw["outlier"]]


def stars(p):
    return "***" if p < 0.001 else "**" if p < 0.01 else "*" if p < 0.05 else "ns"


stat_rows = []
for rep in REPEATS:
    a = clean[(clean["repeat"] == rep) & (clean["time"] == "3h")]["count"].values
    b = clean[(clean["repeat"] == rep) & (clean["time"] == "5h")]["count"].values
    _, p_mw = stats.mannwhitneyu(a, b, alternative="two-sided")
    _, p_t = stats.ttest_ind(a, b, equal_var=False)
    stat_rows.append({"group": rep, "n_3h": len(a), "n_5h": len(b),
                      "mean_3h": a.mean(), "mean_5h": b.mean(),
                      "p_mannwhitney": p_mw, "p_welch_t": p_t, "sig": stars(p_mw)})
a = clean[clean["time"] == "3h"]["count"].values
b = clean[clean["time"] == "5h"]["count"].values
_, p_pool = stats.mannwhitneyu(a, b, alternative="two-sided")
_, p_pool_t = stats.ttest_ind(a, b, equal_var=False)
stat_rows.append({"group": "pooled (3 repeats)", "n_3h": len(a), "n_5h": len(b),
                  "mean_3h": a.mean(), "mean_5h": b.mean(),
                  "p_mannwhitney": p_pool, "p_welch_t": p_pool_t, "sig": stars(p_pool)})
stat_df = pd.DataFrame(stat_rows)
stat_df.to_csv(os.path.join(OUT, "ss_control_stats.csv"), index=False, encoding="utf-8-sig")

from docx import Document
from docx.shared import Pt, Inches

doc = Document()
doc.add_heading("SS-control LIVE/DEAD Re-count Report (ImageJ)", 0)
p = doc.add_paragraph()
p.add_run("Date: %s" % datetime.date.today().isoformat()).italic = True

doc.add_heading("Sample & methods", level=1)
doc.add_paragraph(
    "Sample: stainless-steel control coupons (ss-control) from the LIVE/DEAD bacterial "
    "adhesion assay. Single green stain (SYTO9 only, no PI), i.e. all attached bacteria "
    "are stained - the counts below are TOTAL attached bacteria, not live/dead split.")
doc.add_paragraph(
    "Counting engine: ImageJ 1.54 (classic, headless -batch) with adaptive threshold "
    "(background histogram mode + 40) and Analyze Particles (size 3-500 px, 8-connected). "
    "Images are single-channel greyscale TIFFs named {n}.tif with signal in the green "
    "channel; each image was fed to the macro as {n}-g.tif so the macro selects the green "
    "channel, and results were mapped back to the original names. "
    "Field area (40x objective): 46946.10 um2, so density [cells/cm2] = count / 46946.10 x 1e8. "
    "Outlier screening: within each repeat x time group, fields with |x - median| > 5 x MAD "
    "(median absolute deviation) were flagged and excluded from the summary and statistics.")
doc.add_paragraph(
    "Repeats analysed: %s (independent repeats of the same assay)." % ", ".join(REPEATS))
doc.add_paragraph(
    "Note on completeness: repeat 2 contains 16 fields at 3 h and only 12 fields at 5 h "
    "(repeat 1 and repeat 3 have 15 + 15). Exposure times also differ between repeats "
    "(repeat 1: 1.0-1.8 s; repeat 2: 0.6-1.8 s; repeat 3: 1.0 s) - the adaptive threshold "
    "normalises for this, but it is worth knowing when comparing repeats.")

doc.add_heading("Summary by repeat (counts per field, mean +/- SD)", level=1)
t = doc.add_table(rows=1, cols=len(by_rep.columns))
t.style = "Light Grid Accent 1"
for j, c in enumerate(by_rep.columns):
    t.rows[0].cells[j].text = str(c)
for _, r in by_rep.iterrows():
    cells = t.add_row().cells
    for j, c in enumerate(by_rep.columns):
        v = r[c]
        cells[j].text = ("%.1f" % v) if isinstance(v, float) else str(v)

doc.add_heading("Summary, 3 repeats pooled", level=1)
t2 = doc.add_table(rows=1, cols=len(summ.columns))
t2.style = "Light Grid Accent 1"
for j, c in enumerate(summ.columns):
    t2.rows[0].cells[j].text = str(c)
for _, r in summ.iterrows():
    cells = t2.add_row().cells
    for j, c in enumerate(summ.columns):
        v = r[c]
        cells[j].text = ("%.1f" % v) if isinstance(v, float) else str(v)

doc.add_heading("3 h vs 5 h statistics", level=1)
t3 = doc.add_table(rows=1, cols=len(stat_df.columns))
t3.style = "Light Grid Accent 1"
for j, c in enumerate(stat_df.columns):
    t3.rows[0].cells[j].text = str(c)
for _, r in stat_df.iterrows():
    cells = t3.add_row().cells
    for j, c in enumerate(stat_df.columns):
        v = r[c]
        cells[j].text = ("%.4g" % v) if isinstance(v, float) else str(v)
doc.add_paragraph(
    "Per-repeat significance was tested with Mann-Whitney U (non-parametric) and reported "
    "alongside Welch's t-test. Note that the direction of the 3 h -> 5 h change is NOT "
    "consistent between repeats (repeat 1 decreases, repeat 2 increases, repeat 3 unchanged), "
    "and the pooled comparison is not significant (p = %.3f)." % p_pool)

doc.add_heading("Figures", level=1)
for fn, cap in (("ss_control_bar.png", "Figure 1. Total attached bacteria, 3 repeats pooled, "
                                       "mean +/- SD with individual fields overlaid."),
                ("ss_control_bar_by_repeat.png", "Figure 2. Total attached bacteria split by "
                                                 "repeat (mean +/- SD with individual fields).")):
    fp = os.path.join(OUT, fn)
    if os.path.exists(fp):
        doc.add_picture(fp, width=Inches(5.8))
        doc.add_paragraph(cap)

doc.add_heading("Output files", level=1)
for fn in ("ss_control_raw.csv", "ss_control_summary.csv",
           "ss_control_summary_by_repeat.csv", "ss_control_stats.csv",
           "ss_control_bar.png", "ss_control_bar_by_repeat.png"):
    doc.add_paragraph(os.path.join(OUT, fn))
doc.add_paragraph("QC images (blue circles + numbering on every counted particle): "
                  + os.path.join(OUT, "qc", "<repeat>", "<time>"))

report = os.path.join(OUT, "ss_control_report.docx")
doc.save(report)
print("报告:", report)
print(stat_df.round(4).to_string(index=False))
