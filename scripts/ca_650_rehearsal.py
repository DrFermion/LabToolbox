# -*- coding: utf-8 -*-
"""End-to-end rehearsal: simulate the user typing a new day into the entry area,
then check that the pipeline picks it up and redraws everything (on a throwaway copy)."""
import importlib.util
import os
import shutil

import openpyxl

SRC = r"E:\LabToolbox\scripts\ca_650_update.py"
TMP = r"C:\Users\PC\AppData\Local\Temp\ca650_rehearsal"
shutil.rmtree(TMP, ignore_errors=True)
os.makedirs(TMP, exist_ok=True)

spec = importlib.util.spec_from_file_location("ca650", SRC)
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)

shutil.copy2(m.XLSX, os.path.join(TMP, "650nm.xlsx"))
m.XLSX = os.path.join(TMP, "650nm.xlsx")
m.BASE = TMP
print("演练副本:", m.XLSX)

# --- simulate: user types "Day 3" into the pale-yellow entry block (r17-r19) ---
wb = openpyxl.load_workbook(m.XLSX)
sim = {"area1": 61.0, "area1-R90": 56.0, "area2": 54.0, "control": 66.0}
for g, v in sim.items():
    ws = wb[g]
    for i, d in enumerate((round(v - 2, 1), round(v, 1), round(v + 2, 1))):
        r = 17 + i
        ws.cell(r, 1, "Day 3")
        ws.cell(r, 2, i + 1)
        ws.cell(r, 3, d)
        ws.cell(r, 4, round(d + 0.8, 1))
        ws.cell(r, 5, round(d + 0.4, 1))
wb.save(m.XLSX)
print("已模拟手填 Day 3 (area1 59/61/63, control 64/66/68 …)")

# --- run the pipeline ---
m.main()

# --- check ---
print()
print("=== 演练校验 ===")
wv = openpyxl.load_workbook(m.XLSX, data_only=True)
ws = wv["summary"]
for r in ws.iter_rows(min_row=3, max_row=ws.max_row):
    vals = [("" if c.value is None else str(c.value)) for c in r[:8]]
    while vals and vals[-1] == "":
        vals.pop()
    if vals:
        print("  summary r%-2d| %s" % (r[0].row, " | ".join(vals)))
for g in ("area1", "control"):
    wsv = wv[g]
    data_last = max([r for r in range(3, wsv.max_row + 1)
                     if any(wsv.cell(r, c).value is not None for c in (1, 2, 3, 4, 5))])
    print(f"  {g}: 数据到 r{data_last} (应含 Day 3), 末行 {wsv.max_row}")
print("  产物:", sorted(f for f in os.listdir(TMP) if f.endswith(".png")))
